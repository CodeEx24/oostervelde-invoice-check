#!/usr/bin/env python3
"""Generates the delivery blueprint: workbook (.xlsx), Jira CSV and validates the plan (sums, refs)."""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.dirname(os.path.abspath(__file__))
KEY = "OIC"

EPICS = [
 ("EP-01","Foundations & dataset","Repo, Dutch number/date handling, dataset loaders, translated contract.","Everything downstream depends on correct Dutch parsing and a translated contract.","P0",1,[],"Contract Art. 4 translated; money/date helpers tested; dataset loads."),
 ("EP-02","Invoice parsing (4 layouts)","Normalise Hoekstra, De Wit, Bosman and Jansen invoices into one schema.","Without a clean schema no rule can be trusted.","P0",1,["EP-01"],"All 40 invoices parse; every parsed invoice reconciles with its own stated totals unless the invoice itself is wrong."),
 ("EP-03","Contract rule engine","One deterministic rule per Article 4 clause; findings with article, line, amount, evidence.","Amounts must be correct to the cent; Sanne must be able to forward the reasoning unedited.","P0",2,["EP-02"],"Golden table for all 40 invoices passes; no false flags on clean invoices."),
 ("EP-04","MCP server","Small, clearly named tools over the dataset and engine.","The skill needs tools, not prompts, for facts and maths.","P0",3,["EP-03"],"11 tools callable over stdio by a real MCP client."),
 ("EP-05","Skill & Sanne-in-the-loop","Skill that presents, asks Sanne, logs decisions and overrides; never sends.","Sanne signs off every payment; trust is the product.","P0",3,["EP-04"],"Skill covers present → decide → log → draft → close; no send path exists."),
 ("EP-06","The run, comparison & drafts","Run all 40, total the dispute, compare with Sanne's 10, Dutch e-mail drafts.","The assignment's measurable output.","P1",3,["EP-03"],"Review list, comparison and drafts regenerate with one command."),
 ("EP-07","Verification & submission","Independent check, NOTES.md, README, video, submit.","Explainability and honesty are scored.","P1",4,["EP-06"],"NOTES.md has the four headings; README runs in 10 minutes; video recorded; mail sent."),
]
FEATURES = [
 ("F-01","EP-01","Dutch formats & loaders","Money, dates, CSV, contract translation","P0","M"),
 ("F-02","EP-02","Supplier parsers","Four layouts to one schema + self-check","P0","L"),
 ("F-03","EP-03","Price & indexation","Execution-date pricing, unit conversion","P0","L"),
 ("F-04","EP-03","Work-order rules","Call-out fee, hours/MU, meerwerk/GA, unknown WO","P0","L"),
 ("F-05","EP-03","VAT rule","9 % / 21 % by code and build year","P0","M"),
 ("F-06","EP-03","Cross-invoice rules","Duplicates and credit-note netting","P0","M"),
 ("F-07","EP-03","Arithmetic & verdict","Own-arithmetic checks, exact VAT, verdict logic","P0","M"),
 ("F-08","EP-04","MCP tools","Read tools, check, review_all, drafts, decision log","P0","L"),
 ("F-09","EP-05","Skill & HITL","Skill workflow and override loop","P0","M"),
 ("F-10","EP-06","Run & compare","Review list, comparison, Dutch drafts, golden tests","P1","M"),
 ("F-11","EP-07","Verification & delivery","Blind audit, NOTES, README, video, submit","P1","M"),
]
# id, epic, feature, summary, user story, AC list, prio, pts, wave, ws, sprint, deps, status
S = []
def st(*a): S.append(a)
st("US-001","EP-01","F-01","Translate Article 4 and build a glossary","As the builder I want the contract in English so I don't misread Dutch rules.",["Every Art. 4.1-4.10 clause translated","Glossary of invoice terms (spoed, meerwerk, voorrijkosten, ...)","Dutch text kept verbatim for quoting"],"P0",3,1,"WS-A",1,[],"Done")
st("US-002","EP-01","F-01","Scaffold TypeScript/vitest/tsx project","As the builder I want a runnable repo in minutes.",["npm install && npm test works from clean checkout","Strict TypeScript"],"P0",2,1,"WS-A",1,[],"Done")
st("US-003","EP-01","F-01","Dutch money and date parsing","As the builder I want 1.234,56 and 02/01/2026 parsed correctly.",["Money in integer cents, no floats","DD/MM, DD-MM, ISO and '5 januari 2026' all parse; impossible dates rejected"],"P0",3,1,"WS-A",1,["US-002"],"Done")
st("US-004","EP-01","F-01","Load price list, work orders, contract text, Sanne's verdicts","As the engine I want typed reference data.",["Quoted CSV fields handled","DATA_DIR override works"],"P0",2,1,"WS-A",1,["US-002"],"Done")
st("US-005","EP-02","F-02","Hoekstra parser","As Sanne I want Hoekstra invoices read correctly, including credit notes.",["Lines, units (incl. 'rol 25 m', 'post'), totals, VAT parsed","Credit note links to original invoice"],"P0",3,1,"WS-A",1,["US-003"],"Done")
st("US-006","EP-02","F-02","Bosman parser","As Sanne I want Bosman invoices read correctly.",["Pipe layout parsed","ISO dates; meerwerk flagged"],"P0",2,1,"WS-A",1,["US-003"],"Done")
st("US-007","EP-02","F-02","Jansen parser (day-first dates)","As Sanne I want Jansen invoices read correctly.",["Two-line item layout parsed","02/01/2026 = 2 January, not 1 February"],"P0",3,1,"WS-A",1,["US-003"],"Done")
st("US-008","EP-02","F-02","De Wit parser (letter style)","As Sanne I want De Wit's letter-style invoices read correctly.",["Dutch month names","VAT rate per line","Meerwerk line containing an extra colon"],"P0",5,1,"WS-A",1,["US-003"],"Done")
st("US-009","EP-02","F-02","Parser self-check against stated totals","As the builder I want parser errors separated from invoice errors.",["Lines sum, VAT and total checked per invoice","Only 2026-0052 and DW26-016 have real arithmetic errors"],"P0",2,1,"WS-A",1,["US-005","US-006","US-007","US-008"],"Done")
st("US-010","EP-03","F-03","Price and indexation by execution date, with unit conversion","As Sanne I want prices checked at the right year's price.",["Work executed before 2026-01-01 uses 2025 price; else +3.2 % rounded per unit","rol 25 m vs m compared per list unit","Invoice date never decides the price year"],"P0",5,2,"WS-B",2,["US-003","US-004","US-009"],"Done")
st("US-011","EP-03","F-04","Call-out fee rule (Art. 4.3)","As Sanne I want call-out fees only on spoed and once.",["Regular work: whole fee disputed","Spoed: second fee disputed"],"P0",3,2,"WS-B",2,["US-010"],"Done")
st("US-012","EP-03","F-04","Hours rule with MU approvals (Art. 4.4)","As Sanne I want hours above the maximum caught.",["Max + approved MU hours","Excess valued at contract rate"],"P0",3,2,"WS-B",2,["US-010"],"Done")
st("US-013","EP-03","F-04","Meerwerk rule with GA approvals (Art. 4.5)","As Sanne I want extra work handled exactly.",["<= 250 free","Above 250 without GA: nothing paid (not even 250)","With GA: at most approved amount"],"P0",5,2,"WS-B",2,["US-010"],"Done")
st("US-014","EP-03","F-05","VAT 9 % / 21 % rule (Art. 4.6)","As Sanne I want VAT checked per line, not by counting VAT lines.",["9 % only for W01/W04 on homes older than two years","Overcharged VAT disputed; undercharged VAT held for Sanne"],"P0",5,2,"WS-B",2,["US-010"],"Done")
st("US-015","EP-03","F-04","Off-list items, unknown work orders, supplier mismatch (Art. 4.7, 4.9)","As Sanne I want unknown orders on hold.",["Unknown WO = hold, nothing else claimed","Off-list item flagged as judgment"],"P0",3,2,"WS-B",2,["US-004"],"Done")
st("US-016","EP-03","F-06","Duplicates and credit-note netting (Art. 4.8)","As Sanne I want cross-invoice problems found.",["Identical second invoice for a WO disputed in full","Credit note netted; not disputed twice"],"P0",5,3,"WS-B",3,["US-011","US-012"],"Done")
st("US-017","EP-03","F-07","Invoice arithmetic checks","As Sanne I want the contractor's own sums verified.",["Line = qty x price","Total = subtotal + VAT","Overstated totals shown in their own column"],"P0",3,3,"WS-B",3,["US-009"],"Done")
st("US-018","EP-03","F-07","Exact VAT reclaim per VAT bucket","As Sanne I want VAT correct to the cent.",["Reclaim = billed VAT - VAT on accepted amounts per rate","Per-finding VAT reconciles to invoice-level figure"],"P0",3,3,"WS-B",3,["US-014"],"Done")
st("US-019","EP-03","F-07","Verdict logic and needs-Sanne flag","As Sanne I want to see what is uncertain first.",["approve / dispute / hold rules documented","judgment findings flagged"],"P0",3,3,"WS-B",3,["US-016","US-017","US-018"],"Done")
st("US-020","EP-04","F-08","MCP server skeleton and read tools","As Claude I want tools for facts.",["list_invoices, get_invoice, get_price, get_work_order, get_contract_rule","Unknown WO returns an error mentioning Art. 4.9"],"P0",3,3,"WS-C",3,["US-004","US-009"],"Done")
st("US-021","EP-04","F-08","check_invoice, find_duplicates, review_all","As Claude I want the engine behind tools.",["Findings carry article, line, amount, evidence, Dutch sentence","Errors are clean, not crashes"],"P0",3,3,"WS-C",3,["US-019","US-020"],"Done")
st("US-022","EP-04","F-08","draft_dispute_email (Dutch draft)","As Sanne I want the e-mail written for me.",["Dutch; cites article and amount; 10-day and 30-day wording","Marked CONCEPT; nothing is sent"],"P0",3,3,"WS-C",3,["US-021"],"Done")
st("US-023","EP-04","F-08","Decision log tools","As Sanne I want my decisions and overrides remembered.",["Override requires her reason","Log is append-only JSONL"],"P0",3,3,"WS-C",3,["US-021"],"Done")
st("US-024","EP-04","F-08","Server test with a real MCP client over stdio","As the builder I want proof the server works.",["Tools list, get_price, check_invoice, decisions tested via SDK client"],"P0",3,3,"WS-C",3,["US-023"],"Done")
st("US-025","EP-05","F-09","Skill: present, decide, log, draft, close","As Sanne I want Claude to keep me in charge.",["Judgment items first","Never sends or contacts anyone","Overrides logged with reason"],"P0",3,4,"WS-C",4,["US-021","US-023"],"Done")
st("US-026","EP-05","F-09","Run the skill live in Claude Code/Desktop as Sanne","As the builder I want to see it work end to end.",["Skill triggers on a natural request","All tools reachable","One override logged"],"P0",3,4,"WS-C",4,["US-025"],"Ready")
st("US-027","EP-05","F-09","Override-to-rule-update proposal loop","As the team I want overrides to improve the rules.",["Skill proposes a change; a person edits code","No automatic rule changes"],"P2",2,4,"WS-C",4,["US-025"],"Draft")
st("US-028","EP-06","F-10","Run all 40 and write the review list","As Sanne I want one list with a total.",["review-list.md/.csv/.json generated by one command","Totals: approve/dispute/hold, excl./incl. VAT"],"P0",3,3,"WS-D",3,["US-019"],"Done")
st("US-029","EP-06","F-10","Comparison with Sanne's 10","As the assessor I want every difference explained.",["Table + commentary per invoice","Contract-based, not assuming Sanne is right"],"P0",3,4,"WS-D",4,["US-028"],"Done")
st("US-030","EP-06","F-10","Golden-table regression tests","As the builder I want to know when a number changes.",["All 40 invoices asserted","Rule-level tests with hand-computed values"],"P0",3,4,"WS-D",4,["US-028"],"Done")
st("US-031","EP-07","F-11","Independent blind audit and reconcile","As the builder I want a second opinion that did not see my code.",["Separate agent audits the 40 invoices blind","Every difference resolved or documented"],"P0",3,4,"WS-D",4,["US-030"],"Done")
st("US-032","EP-07","F-11","NOTES.md (four headings)","As the assessor I want decisions, handover, changes and gaps.",["Decisions and translations; Where Sanne decides; What changed; What isn't finished"],"P0",3,4,"WS-D",4,["US-031"],"Done")
st("US-033","EP-07","F-11","README runnable in 10 minutes","As the assessor I want to run it quickly.",["npm install, npm test, npm run all documented with expected output"],"P0",2,4,"WS-D",4,["US-028"],"Done")
st("US-034","EP-07","F-11","Record the <=5-minute video","As the assessor I want to see it running.",["Skill running; one hard case; one thing it still gets wrong"],"P0",3,4,"WS-D",4,["US-026","US-032"],"Ready")
st("US-035","EP-07","F-11","Submit to nigel@mstr.nl and jason@mstr.nl","As the builder I want to deliver within 48 hours.",["Zip/GitHub link + video link in one mail"],"P0",1,4,"WS-D",4,["US-034"],"Draft")

SPRINT_GOALS = {1:"Foundations + all four parsers",2:"Core rules: price, call-out, hours, meerwerk, VAT, unknown WO",3:"Cross-invoice rules, exactness, MCP server, first full run",4:"Skill live, comparison, verification, NOTES, video, submit"}

# ---------- validation ----------
ids = {s[0] for s in S}
for s in S:
    for d in s[11]: assert d in ids, f"{s[0]} depends on unknown {d}"
    assert s[1] in {e[0] for e in EPICS}
    assert s[7] in (1,2,3,5,8), s
    # dependency-first: a story's wave/sprint must not precede its prerequisites
    for d in s[11]:
        dep = next(x for x in S if x[0]==d)
        assert dep[10] <= s[10], f"{s[0]} (sprint {s[10]}) scheduled before dependency {d} (sprint {dep[10]})"
total_pts = sum(s[7] for s in S)
by_sprint = {n: sum(s[7] for s in S if s[10]==n) for n in (1,2,3,4)}
assert sum(by_sprint.values()) == total_pts
print("stories", len(S), "points", total_pts, "per sprint", by_sprint)

# ---------- workbook ----------
wb = Workbook(); wb.remove(wb.active)
HEAD = PatternFill("solid", fgColor="1A1A1A"); ZEB = PatternFill("solid", fgColor="F7F7F8")
STATUS_FILL = {"Done":"E7F4EA","In Progress":"FBF1DE","Ready":"E8F0FB","Draft":"EEEEEE","Blocked":"FBEBEB"}
thin = Side(style="thin", color="E5E7EB")
def sheet(name, header, rows, widths=None, status_col=None):
    ws = wb.create_sheet(name[:31])
    ws.append(header)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = Font(name="Calibri", bold=True, color="FFFFFF"); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    if rows: ws.auto_filter.ref = ws.dimensions
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Calibri", size=10)
            c.border = Border(bottom=thin)
            if i % 2 == 1: c.fill = ZEB
        if status_col is not None:
            v = row[status_col].value
            if v in STATUS_FILL: row[status_col].fill = PatternFill("solid", fgColor=STATUS_FILL[v])
    for i, h in enumerate(header, start=1):
        w = (widths[i-1] if widths and i-1 < len(widths) else 22)
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

sheet("README",["Item","Content"],[
 ["Purpose","Delivery blueprint for the Oostervelde invoice-check build (MCP server + Claude skill), 48-hour solo delivery."],
 ["How to use","Product Backlog = source of truth. PM Import = same data as the Jira CSV. Statuses reflect the real state at generation time."],
 ["Status legend","Draft → Ready → In Progress → Blocked → Code Review → QA Testing → UAT → Done"],
 ["Priority legend","P0 critical path · P1 high · P2 medium · P3 nice-to-have (Jira: Highest/High/Medium/Low)"],
 ["Waves","W1 Foundations & parsing · W2 Core rules · W3 Cross-invoice rules, MCP, run · W4 Skill live, verification, delivery"],
 ["Workstreams","WS-A Data & parsing · WS-B Rule engine · WS-C MCP & skill · WS-D Verification & docs"],
 ["Conventions","Branch feat/oic-<n>/<title> · Conventional Commits with (OIC-###) · see CLAUDE.md"],
 ["Totals",f"{len(S)} stories · {total_pts} points · sprints (points): "+", ".join(f"S{n}={p}" for n,p in by_sprint.items())],
],[22,110])
sheet("Project Overview",["Topic","Detail"],[
 ["Summary","Local MCP server + Claude skill that checks 40 Dutch contractor invoices (4 layouts) against Article 4 of the Woonkring Oostervelde framework contract and gives Sanne a review list; she decides everything."],
 ["Objectives","Correct to the cent · no false flags · evidence Sanne can forward unedited · explain every difference with her 10 verdicts · honest about gaps."],
 ["Target users","Sanne Verbeek, accounts payable (signs off every payment). Assessors at MSTR."],
 ["KPIs","All 40 invoices reviewed · golden table passes · differences with Sanne explained 10/10 · README runs in 10 min · video <= 5 min."],
 ["Constraints","48 hours from receipt · solo · Dutch documents · money in cents · nothing may reach a contractor without Sanne."],
 ["Assumptions","See Gaps & Assumptions."],
 ["Top risks","Misreading Dutch contract rules · parser trusting a wrong layout · overconfident flags · skill not exercised live."],
],[22,110])
sheet("Product Roadmap",["Phase","Wave","Scope","Release target","Priority","Dependencies"],[
 ["MVP / Phase 1","W1","Foundations + 4 parsers","Hour 12","P0","-"],
 ["MVP / Phase 1","W2","Core rules","Hour 24","P0","W1"],
 ["MVP / Phase 1","W3","Cross-invoice rules, MCP server, first run","Hour 36","P0","W2"],
 ["MVP / Phase 1","W4","Skill live, comparison, verification, NOTES, video, submit","Hour 48","P0","W3"],
 ["Phase 2 (after MSTR call)","-","PDF/OCR input, ERP export, learning from overrides","TBD","P3","MVP"],
],[24,8,60,16,10,16])
sheet("Epic Backlog",["Epic ID","Name","Description","Business value","Priority","Wave","Dependencies","Acceptance criteria","Status"],
 [[e[0],e[1],e[2],e[3],e[4],f"W{e[5]}",", ".join(e[6]) or "-",e[7],
   ("Done" if all(s[12]=="Done" for s in S if s[1]==e[0]) else "In Progress")] for e in EPICS],[9,28,44,40,9,7,14,48,12],8)
sheet("Feature Backlog",["Feature ID","Epic","Name","Description","Priority","Complexity","Dependencies","Status"],
 [[f[0],f[1],f[2],f[3],f[4],f[5],", ".join(sorted({d for s in S if s[2]==f[0] for d in s[11] if next(x for x in S if x[0]==d)[2]!=f[0]})) or "-",
   ("Done" if all(s[12]=="Done" for s in S if s[2]==f[0]) else "In Progress")] for f in FEATURES],[10,9,28,46,9,11,24,12],7)
sheet("Product Backlog",["Story ID","Epic","Feature","Summary","User story","Acceptance criteria","Priority","Points","Wave","Workstream","Sprint","Dependencies","Status"],
 [[s[0],s[1],s[2],s[3],s[4],"\n".join("☐ "+a for a in s[5]),s[6],s[7],f"W{s[8]}",s[9],f"S{s[10]}",", ".join(s[11]) or "-",s[12]] for s in S],
 [9,8,8,36,44,52,8,8,6,11,7,18,12],12)
sprint_rows = []
for n in (1,2,3,4):
    ss=[s for s in S if s[10]==n]
    sprint_rows.append([f"S{n}",SPRINT_GOALS[n],", ".join(s[0] for s in ss),by_sprint[n],"; ".join(sorted({d for s in ss for d in s[11] if next(x for x in S if x[0]==d)[10]!=n})) or "-","Working increment: "+SPRINT_GOALS[n]])
sheet("Sprint Backlog",["Sprint","Goal","Stories","Est. points","Dependencies (earlier sprints)","Deliverables"],sprint_rows,[8,44,70,10,30,44])
sheet("Technical Backlog",["Area","Task","Notes"],[
 ["Infra","Node 22 + TypeScript strict + tsx + vitest","No build step needed to run"],
 ["Data","Typed loaders for CSV/MD dataset; DATA_DIR override","No database"],
 ["API","11 MCP tools over stdio","See API Contracts"],
 ["Backend","Engine: one block per contract article","Integer cents everywhere"],
 ["Security","Read-only tools; decision log is the only write; no network; no mail integration","Human in the loop by construction"],
 ["DevOps","npm scripts; .mcp.json; Claude Desktop config snippet","See README"],
 ["Monitoring","Tool errors are explicit messages, not crashes","Parsers throw on unknown layout"],
],[14,70,46])
sheet("Architecture & Decisions",["Topic","Decision","Why"],[
 ["Components","parsers → engine → render → (cli | MCP server) → skill","Each step testable alone"],
 ["Stack","TypeScript + MCP SDK + zod","Locked; no swaps without an ADR"],
 ["Rules in code, not prompt","All contract maths in deterministic code","Cent-exact, testable, explainable"],
 ["Execution date","Price year follows execution date (Art. 4.2)","Contract text; Sanne's shortcut is wrong"],
 ["Verdicts","approve / dispute / hold + judgment flag","Matches the brief; uncertainty visible"],
 ["VAT reclaim","Per VAT bucket at invoice level","How invoices round; matches a corrected invoice"],
 ["HITL","Proposals only; decisions logged; overrides need reasons","Sanne signs off"],
 ["AI integration","Claude presents, asks, logs, drafts; never computes","See skill"],
],[24,60,52])
sheet("Data Model",["Entity","Fields","Relationships","Notes"],[
 ["PriceItem","code, supplier, description, unit, price2025Cents","-","From price-agreement-2025.csv"],
 ["WorkOrder","wo, supplier, type, maxLaborHours, MU approval + hours, GA approval + amount, built","1 WO : n Invoice","Export is the authority"],
 ["Invoice","supplier, number, dates (ISO), executionDate, workOrder, isCredit, creditOf, totals","1 : n InvoiceLine, VatLine","Normalised from 4 layouts"],
 ["InvoiceLine","description, qtyMilli, unit, unitPriceCents, amountCents, vatRatePct?, code?, isMeerwerk","n : 1 Invoice","qty in thousandths"],
 ["Finding","article, kind, severity, confidence, amounts (excl/VAT/other), reasons EN/NL, evidence, Dutch quote","n : 1 InvoiceReview","Evidence Sanne can forward"],
 ["InvoiceReview","verdict, findings, billed, disputed, creditApplied, needsSanne","1 : 1 Invoice",""],
 ["Decision","at, invoiceNo, proposed, final, override, reason","n : 1 Invoice","JSONL, append-only"],
 ["Migrations / seed","None (no database)","",""],
],[18,70,26,28])
sheet("API Contracts",["Tool","Request","Response","Auth","Status"],[
 ["list_invoices","{}","[{file, supplier, invoiceNo, dates, workOrder, total}]","local stdio","Done"],
 ["get_invoice","{invoice}","normalised invoice (cents)","local stdio","Done"],
 ["get_price","{code, execution_date}","applicable price, indexed?","local stdio","Done"],
 ["get_work_order","{work_order}","work order or error (Art. 4.9)","local stdio","Done"],
 ["get_contract_rule","{query?}","[{id, titleNl, textNl, textEn}]","local stdio","Done"],
 ["check_invoice","{invoice}","verdict + findings + amounts + evidence + Dutch text","local stdio","Done"],
 ["find_duplicates","{}","duplicates, repeats, credit notes","local stdio","Done"],
 ["review_all","{format?}","markdown or json review list + totals","local stdio","Done"],
 ["draft_dispute_email","{invoice}","Dutch draft text (never sent)","local stdio","Done"],
 ["record_decision","{invoice, final_verdict, sanne_reason, final_dispute_excl_eur?}","logged decision (override flag)","local stdio","Done"],
 ["list_decisions","{}","decisions + overrides","local stdio","Done"],
],[22,52,48,14,10],4)
sheet("QA & Testing",["Feature","Unit","Integration","Edge cases","Regression"],[
 ["Dutch formats","money.test: parse/format/rounding/indexation of all 19 prices","-","31-02, '5 smarch', 02/01 vs 01/02","parsers.test"],
 ["Parsers","per layout","all 40 parse; arithmetic sweep finds exactly 2 real errors","'post' unit, meerwerk with extra colon, credit note","parsers.test"],
 ["Rules","per article with hand-computed values","golden table for 40 invoices","spoed twice, MU hours, GA cap, 9 % on new build","engine.test (40 golden invoices + rule tests)"],
 ["MCP server","-","real MCP client over stdio","unknown invoice, unknown WO, override without reason","server.test"],
 ["Skill","-","live run (US-026, open)","Sanne overrides a proposal","manual"],
 ["UAT","-","Sanne persona walk-through","-","-"],
],[18,44,44,44,24])
sheet("DevOps & Deploy",["Topic","Detail"],[
 ["Repo","Single repo; trunk-based; branch feat/oic-<n>/<title>"],
 ["CI","npm ci && npm run typecheck && npm test"],
 ["Environments","Local only (stdio MCP server); DATA_DIR/OUTPUT_DIR/DECISIONS_FILE env vars"],
 ["Logging","Tool errors returned as messages; decisions in output/decisions.jsonl"],
 ["Backup / DR","Dataset is read-only; outputs regenerate with `npm run all`"],
],[22,100])
sheet("Risk Register",["Risk","Impact","Probability","Mitigation","Owner"],[
 ["Contract misread (Dutch)","High","Medium","Quote Dutch verbatim, translate in NOTES, blind second audit","Builder"],
 ["Overconfident false flag","High","Low","needsSanne flag for judgment; tests on clean invoices","Builder"],
 ["Parser trusts wrong layout","High","Low","Self-check vs stated totals; throw on unknown layout","Builder"],
 ["Skill not exercised live","Medium","Medium","US-026 before recording video","Builder"],
 ["Ambiguity resolved differently by MSTR key","Medium","Medium","Documented per decision; amounts shown both ways where small","Builder"],
 ["Missing received-date for 10-day deadline","Low","High","Counted from invoice date; stated in list","Builder"],
],[34,10,12,60,12])
sheet("Future Improvements",["Area","Idea"],[
 ["Input","PDF/OCR invoices; unseen layouts via a learned parser with human confirmation"],
 ["Integration","ERP/booking export of approved lines; mail integration with explicit per-message approval"],
 ["Learning","Turn repeated Sanne overrides into proposed rule changes (human-reviewed)"],
 ["Contract","Load full contract, not just Art. 4; Dutch public-holiday calendar for deadlines"],
 ["Scale","Multi-month history for cumulative hours/meerwerk per work order"],
],[16,110])
# PM import
pm_cols = ["Issue ID","Issue Type","Summary","Description","Epic Name","Epic Link","Priority","Story Points","Status","Labels","Sprint","Wave","Workstream","Blocked By"]
PMAP={"P0":"Highest","P1":"High","P2":"Medium","P3":"Low"}
pm_rows=[]
for e in EPICS:
    pm_rows.append([e[0],"Epic",e[1],f"{e[2]}\n\nBusiness value: {e[3]}\nAcceptance: {e[7]}\nWave W{e[5]}",e[1],"",PMAP[e[4]],"", "Done" if all(s[12]=="Done" for s in S if s[1]==e[0]) else "In Progress",f"wave-{e[5]} epic-{e[0].lower()} phase-mvp {e[4].lower()}","",f"W{e[5]}","",""])
for s in S:
    ep = next(e for e in EPICS if e[0]==s[1])
    desc = f"{s[4]}\n\nAcceptance criteria:\n"+"\n".join("- [ ] "+a for a in s[5])+f"\n\nDone gates: tests green, no regressions, reviewed PR, docs updated.\nMapping: Epic {s[1]} · Feature {s[2]} · Wave W{s[8]} · {s[9]} · Sprint S{s[10]}"
    pm_rows.append([s[0],"Story",s[3],desc,"",ep[1],PMAP[s[6]],s[7],s[12],f"wave-{s[8]} {s[9].lower()} sprint-{s[10]} epic-{s[1].lower()} phase-mvp {s[6].lower()}",f"S{s[10]}",f"W{s[8]}",s[9],", ".join(s[11])])
sheet("PM Import",pm_cols,pm_rows,[10,9,40,60,24,26,10,8,12,44,7,6,10,22],8)
with open(os.path.join(OUT,"jira-import.csv"),"w",newline="",encoding="utf-8") as f:
    w=csv.writer(f); w.writerow(pm_cols); w.writerows(pm_rows)
sheet("Agile & AI Practices",["Topic","Practice"],[
 ["Definition of Ready","AC testable · prerequisites done · <= 8 pts · design/contract linked · human-in-the-loop implications noted"],
 ["Definition of Done","AC met · reviewed PR · npm test green · golden table updated with a reason if numbers change · NOTES.md updated if a decision changed · no float money"],
 ["Ceremonies","Solo: 4 timeboxes (~12 h), demo at the end of each (green tests + regenerated review list)"],
 ["AI-assisted dev","Claude builds from the translated contract; every rule needs a hand-computed test"],
 ["AI-assisted test","Independent blind audit by a separate agent that never sees the code; differences reconciled"],
 ["AI-assisted review","Findings must cite the article and quote the Dutch text; no claim without a tool result"],
 ["AI in the product","Claude presents, asks and logs; code decides numbers; Sanne decides outcomes"],
],[24,110])
sheet("Gaps & Assumptions",["#","Gap / assumption","Impact","Recommendation / what we did"],[
 [1,"Date invoice was RECEIVED is not in the data (10-working-day deadline, 30-day payment)","Low","Count from invoice date; say so in the list. Ask MSTR/Sanne for the received date."],
 [2,"Unit conversion rounding (167,70 / 25 = 6,708 per m)","Low (8 cents on h-2026-0057)","Compare exactly (no intermediate rounding): 21,68. If rounded to 6,71: 21,60. Ask MSTR."],
 [3,"Contractor UNDER-charges VAT (9 % on new build / non-painting)","Medium","No money to reclaim; the invoice is wrong → HOLD for corrected invoice. Ask MSTR whether they expect dispute."],
 [4,"Off-list item (Vogelwering): prior approval may exist outside the work-order export","Medium","Dispute, flagged judgment: Sanne confirms."],
 [5,"'Older than two years' uses build year vs execution year (> 2)","Low","Data has only 1968-1972 and 2025; no border case."],
 [6,"Hours/meerwerk are aggregated per work order only across invoices in this batch","Low","Earlier months not available."],
 [7,"No holiday calendar for working days","Low","Mon-Fri only."],
 [8,"The skill was not yet run live in Claude by the time of generation","Medium","US-026 before the video."],
],[5,70,24,70])

wb.save(os.path.join(OUT,"Oostervelde-Delivery-Blueprint.xlsx"))
print("workbook tabs:", wb.sheetnames, len(wb.sheetnames))
